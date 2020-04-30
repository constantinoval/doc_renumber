from renumber_figures import analize_figures
from renumber_formulas import analize_formulas
from renumber_tables import analize_tables
from docx import Document
from colorama import Fore, init
from pprint import pprint


def doc_analysis(doc):
    init(autoreset=True)
    f = analize_figures(Document(doc))
    e = analize_formulas(Document(doc))
    t = analize_tables(Document(doc))
    print(Fore.GREEN + 'Figures:', len(f))
    pprint(list(f.keys()), indent=5, width=40, compact=True)
    print(Fore.GREEN + 'Equations:', len(e))
    pprint(list(e.keys()), indent=5, width=40, compact=True)
    print(Fore.GREEN + 'Tables:', len(t))
    pprint(list(t.keys()), indent=5, width=40, compact=True)
